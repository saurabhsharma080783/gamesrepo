import json

import docx
import pptx
import pytest

from azure_assessment.analysis.assessor import assess
from azure_assessment.cli import FORMATS, generate_reports, main
from azure_assessment import loaders
from azure_assessment.loaders import csv_loader, json_loader, sample
from azure_assessment.loaders.normalize import to_resource


def _res(type_, **kw):
    return to_resource({"id": f"/subscriptions/s1/resourceGroups/rg/providers/{type_}/x",
                        "name": "x", "type": type_, **kw})


def _rule_ids(resources, cfg=None):
    from azure_assessment.models import Inventory
    return {f.rule_id for f in assess(Inventory(resources), cfg).findings}


def test_storage_rules():
    bad = _res("Microsoft.Storage/storageAccounts", sku={"name": "Standard_LRS"}, tags={},
               properties={"supportsHttpsTrafficOnly": False, "allowBlobPublicAccess": True,
                           "minimumTlsVersion": "TLS1_0"})
    assert {"SEC-001", "SEC-002", "SEC-003", "REL-003", "GOV-001"} <= _rule_ids([bad])
    good = _res("Microsoft.Storage/storageAccounts", sku={"name": "Standard_ZRS"},
                tags={"environment": "p", "owner": "o", "costCenter": "c"},
                properties={"supportsHttpsTrafficOnly": True, "allowBlobPublicAccess": False,
                            "minimumTlsVersion": "TLS1_2"})
    assert _rule_ids([good]) == set()


@pytest.mark.parametrize("rule,expected", [
    ({"sourceAddressPrefix": "*", "destinationPortRange": "3389"}, True),
    ({"sourceAddressPrefix": "Internet", "destinationPortRange": "20-25"}, True),
    ({"sourceAddressPrefix": "0.0.0.0/0", "destinationPortRanges": ["443", "22"]}, True),
    ({"sourceAddressPrefix": "Internet", "destinationPortRange": "443"}, False),
    ({"sourceAddressPrefix": "10.0.0.0/8", "destinationPortRange": "22"}, False),
])
def test_nsg_management_ports(rule, expected):
    nsg = _res("Microsoft.Network/networkSecurityGroups", properties={"securityRules": [
        {"name": "r", "properties": {"direction": "Inbound", "access": "Allow", **rule}}]})
    assert ("SEC-007" in _rule_ids([nsg])) is expected


def test_vm_power_state_and_region():
    vm = _res("Microsoft.Compute/virtualMachines", location="brazilsouth",
              properties={"extended": {"instanceView": {"powerState": {"code": "PowerState/stopped"}}}})
    ids = _rule_ids([vm], {"allowed_locations": ["eastus"]})
    assert {"COST-003", "REL-001", "GOV-002"} <= ids
    assert "GOV-002" not in _rule_ids([vm])  # no allowed_locations configured -> rule skipped


def test_scores_and_aggregates():
    a = assess(sample.build(), customer="Contoso")
    assert a.total_resources > 100
    assert 0 <= a.overall_score <= 100
    assert sum(a.severity_counts.values()) == len(a.findings)
    assert all(0 <= p.score <= 100 for p in a.pillar_scores)
    assert a.key_observations()
    assert sum(len(v) for v in a.roadmap().values()) == len(a.rule_summaries)


def test_load_az_graph_output(tmp_path):
    p = tmp_path / "graph.json"
    p.write_text(json.dumps({"count": 1, "data": [{
        "id": "/subscriptions/abc/resourceGroups/RG/providers/Microsoft.Web/sites/app1",
        "name": "app1", "type": "microsoft.web/sites", "location": "eastus",
        "resourceGroup": "RG", "subscriptionId": "abc", "properties": {"httpsOnly": False}}]}))
    inv = loaders.load(p)
    assert inv.resources[0].subscription_id == "abc"
    assert "SEC-006" in {f.rule_id for f in assess(inv).findings}


def test_roundtrip_export(tmp_path):
    inv = sample.build()
    loaded = loaders.load(json_loader.save(inv, tmp_path / "inv.json"))
    assert len(loaded.resources) == len(inv.resources)
    assert loaded.subscriptions == inv.subscriptions


def test_generate_all_reports(tmp_path):
    out = generate_reports(sample.build(), tmp_path, "Contoso Ltd", FORMATS)
    assert set(out) == {"html", "docx", "pptx", "findings"}

    html = out["html"].read_text()
    assert "Contoso Ltd" in html and "Executive summary" in html and "<script" in html

    d = docx.Document(out["docx"])
    headings = [p.text for p in d.paragraphs if p.style.name.startswith("Heading")]
    assert "1. Executive summary" in headings and "7. Remediation roadmap" in headings and "3. Architecture assessment" in headings
    assert len(d.inline_shapes) >= 4  # charts embedded

    prs = pptx.Presentation(out["pptx"])
    assert len(prs.slides) >= 9
    for i, slide in enumerate(prs.slides, 1):
        for sh in slide.shapes:  # nothing may spill off the slide
            assert sh.left >= 0 and sh.top >= 0, (i, sh.name)
            assert sh.left + sh.width <= prs.slide_width, (i, sh.name)
            assert sh.top + sh.height <= prs.slide_height, (i, sh.name)


def test_cli_demo(tmp_path, capsys):
    assert main(["demo", "-o", str(tmp_path), "-f", "html"]) == 0
    assert list(tmp_path.glob("*.html"))
    with pytest.raises(SystemExit):
        main(["demo", "-o", str(tmp_path), "-f", "pdf"])


# ---- customer input files ----------------------------------------------------

def test_csv_template_roundtrip_matches_json(tmp_path):
    """A Resource Graph style CSV (nested columns as JSON strings) assesses exactly like the JSON."""
    inv = sample.build()
    from_csv = loaders.load(csv_loader.save(inv.resources, tmp_path / "inv.csv"))
    assert len(from_csv.resources) == len(inv.resources)
    a_csv, a_ref = assess(from_csv), assess(inv)
    assert a_csv.coverage_gaps == []
    assert [(f.rule_id, f.resource_id) for f in a_csv.findings] == [(f.rule_id, f.resource_id) for f in a_ref.findings]
    assert a_csv.overall_score == a_ref.overall_score


def test_portal_export_csv(tmp_path):
    """Azure portal 'All resources' export: friendly names, no ids, no properties."""
    p = tmp_path / "portal.csv"
    p.write_text('"NAME","TYPE","RESOURCE GROUP","LOCATION","SUBSCRIPTION"\n'
                 '"vm-web-01","Virtual machine","rg-web","East US","Prod"\n'
                 '"stweb01","Storage account","rg-web","West Europe","Prod"\n'
                 '"legacy","Cloud service (classic)","rg-old","East US","Dev"\n', encoding="utf-8-sig")
    inv = loaders.load(p)
    vm = inv.resources[0]
    assert vm.type == "microsoft.compute/virtualmachines" and vm.location == "eastus"
    assert vm.subscription_name == "Prod" and vm.resource_group == "rg-web"
    assert set(inv.subscriptions.values()) == {"Prod", "Dev"}

    a = assess(inv, {"allowed_locations": ["eastus"]})
    ids = {f.rule_id for f in a.findings}
    assert ids == {"GOV-002", "GOV-003"}  # only checks that don't need properties/tags can run
    gap_ids = {g.rule.id for g in a.coverage_gaps}
    assert {"GOV-001", "SEC-001", "REL-001", "COST-003"} <= gap_ids
    scores = {ps.pillar: ps.score for ps in a.pillar_scores}
    assert scores["Security"] is None and scores["Governance"] is not None  # unknown != passed
    assert a.tag_coverage is None and "properties" in a.missing_fields


def test_flattened_columns_and_semicolons(tmp_path):
    p = tmp_path / "flat.csv"
    p.write_text("Resource Name;Resource Type;Resource Group;Region;Subscription ID;tags.owner;"
                 "properties.supportsHttpsTrafficOnly;properties.minimumTlsVersion;SKU\n"
                 "st1;Microsoft.Storage/storageAccounts;rg;eastus;sub1;alice;false;TLS1_2;Standard_LRS\n")
    inv = loaders.load(p)
    r = inv.resources[0]
    assert r.tags == {"owner": "alice"} and r.sku == {"name": "Standard_LRS"}
    assert r.properties == {"supportsHttpsTrafficOnly": False, "minimumTlsVersion": "TLS1_2"}
    ids = {f.rule_id for f in assess(inv).findings}
    assert {"SEC-001", "REL-003", "GOV-001"} <= ids and "SEC-003" not in ids


def test_tag_string_formats():
    assert to_resource({"name": "a", "type": "t", "tags": "env:prod; owner=bob"}).tags == {"env": "prod", "owner": "bob"}
    assert to_resource({"name": "a", "type": "t", "tags": '{"env": "prod"}'}).tags == {"env": "prod"}


def test_merge_multiple_files_dedupes(tmp_path):
    inv = sample.build()
    half = len(inv.resources) // 2
    a = csv_loader.save(inv.resources[: half + 5], tmp_path / "a.csv")
    b = json_loader.save(type(inv)(inv.resources[half:], inv.subscriptions), tmp_path / "b.json")
    merged = loaders.load([a, b])
    assert len(merged.resources) == len(inv.resources)


@pytest.mark.parametrize("content,suffix,msg", [
    ("foo,bar\n1,2\n", ".csv", "name and type"),
    ("{not json", ".json", "not valid JSON"),
    ("x", ".xlsx", "unsupported file type"),
])
def test_bad_inputs(tmp_path, content, suffix, msg):
    p = tmp_path / f"bad{suffix}"
    p.write_text(content)
    with pytest.raises(ValueError, match=msg):
        loaders.load(p)


def test_cli_report_and_template(tmp_path, capsys):
    tpl = tmp_path / "inventory-template.csv"
    assert main(["template", "-o", str(tpl)]) == 0
    assert main(["report", str(tpl), "-c", "Fabrikam", "-o", str(tmp_path / "out"), "-f", "html",
                 "--inventory-date", "2026-09-01"]) == 0
    html = (tmp_path / "out" / "fabrikam-azure-assessment.html").read_text()
    assert "Fabrikam" in html and "2026-09-01" in html
    assert main(["report", str(tmp_path / "missing.csv"), "-o", str(tmp_path)]) == 2
    assert "not found" in capsys.readouterr().err


def test_reports_render_with_coverage_gaps(tmp_path):
    p = tmp_path / "portal.csv"
    p.write_text("NAME,TYPE,RESOURCE GROUP,LOCATION,SUBSCRIPTION\nvm1,Virtual machine,rg,East US,Prod\n")
    out = generate_reports(loaders.load(p), tmp_path / "out", "Tailspin", FORMATS)
    assert "Data coverage" in out["html"].read_text()
    d = docx.Document(out["docx"])
    assert any(p.text == "Data coverage and limitations" for p in d.paragraphs)
    titles = [sh.text_frame.text for sl in pptx.Presentation(out["pptx"]).slides for sh in sl.shapes
              if sh.has_text_frame]
    assert "Data coverage & limitations" in titles



def test_default_formats_are_word_and_html(tmp_path):
    assert main(["demo", "-o", str(tmp_path)]) == 0
    assert sorted(p.suffix for p in tmp_path.iterdir()) == [".docx", ".html", ".json"]


# ---- infrastructure overview (narrative) --------------------------------------

def test_narrative_covers_estate_and_links_findings():
    from azure_assessment.analysis import narrative

    a = assess(sample.build(), {"allowed_locations": ["eastus", "westeurope", "uksouth"]}, "Contoso")
    secs = {s.key: s for s in narrative.build(a)}
    assert {"estate", "subscriptions", "regions", "compute", "network", "storage", "data", "web",
            "security", "monitoring", "governance"} <= set(secs)
    text = " ".join(p for s in secs.values() for p in s.paragraphs + s.bullets)
    assert "152 Azure resources" in text
    assert "stopped but still allocated" in text and "unattached disks" in text
    assert {r.rule.id for r in secs["compute"].related} >= {"REL-001", "COST-001", "COST-003"}
    assert {r.rule.id for r in secs["regions"].related} == {"GOV-002"}
    # every finding raised is referenced from some part of the overview
    referenced = {r.rule.id for s in secs.values() for r in s.related}
    assert {s.rule.id for s in a.rule_summaries} <= referenced


def test_narrative_without_configuration_data(tmp_path):
    from azure_assessment.analysis import narrative

    p = tmp_path / "portal.csv"
    p.write_text("NAME,TYPE,RESOURCE GROUP,LOCATION,SUBSCRIPTION\n"
                 "vm1,Virtual machine,rg,East US,Prod\nd1,Disk,rg,East US,Prod\n"
                 "st1,Storage account,rg,East US,Prod\n")
    secs = {s.key: s for s in narrative.build(assess(loaders.load(p)))}
    text = " ".join(p for s in secs.values() for p in s.paragraphs)
    assert "could not be determined" in text  # says what it can't see instead of guessing
    assert "tags, so tagging could not be reviewed" in text
    assert "0 of" not in secs["storage"].paragraphs[0]


def test_reports_include_infrastructure_overview(tmp_path):
    out = generate_reports(sample.build(), tmp_path, "Contoso")
    html = out["html"].read_text()
    assert 'id="environment"' in html and 'href="#rule-SEC-007"' in html and 'id="rule-SEC-007"' in html
    d = docx.Document(out["docx"])
    heads = [p.text for p in d.paragraphs if p.style.name.startswith("Heading")]
    i3, i4 = heads.index("4. Infrastructure overview"), heads.index("5. Inventory charts and statistics")
    assert i4 - i3 >= 10  # one subsection per area
    body = "\n".join(p.text for p in d.paragraphs)
    assert "(see 6.1 Security)" in body and "Context: see 4." in body


# ---- architecture assessment -------------------------------------------------

def _arch(resources, cfg=None):
    from azure_assessment.analysis import architecture
    from azure_assessment.models import Inventory
    return architecture.analyse(assess(Inventory(resources), cfg))


def _vnet(name, sub="s1", peers=(), subnets=(), loc="eastus"):
    vid = f"/subscriptions/{sub}/resourceGroups/rg/providers/Microsoft.Network/virtualNetworks/{name}"
    return to_resource({"id": vid, "name": name, "type": "Microsoft.Network/virtualNetworks", "location": loc,
                        "properties": {
                            "subnets": [{"name": s} for s in subnets],
                            "virtualNetworkPeerings": [{"properties": {"remoteVirtualNetwork": {
                                "id": f"/subscriptions/{sub}/resourceGroups/rg/providers/Microsoft.Network/"
                                      f"virtualNetworks/{p}"}}} for p in peers]}})


def test_architecture_sample_topology():
    arch = assess_arch = _arch(sample.build().resources)
    t = arch.topology
    assert t.pattern == "Multi-region hub-and-spoke with exceptions" and t.confidence == "High"
    names = lambda ids: sorted(t.nodes[i].name for i in ids)  # noqa: E731
    assert names(t.hubs) == ["vnet-hub-eastus", "vnet-hub-westeurope"]
    assert names(t.spokes) == ["vnet-dev-eastus", "vnet-prd-eastus", "vnet-prd-westeurope"]
    assert names(t.isolated) == ["vnet-dev-sandbox"]
    assert {e[2] for e in t.edges} == {"hub-spoke", "hub-hub", "spoke-spoke"}
    assert t.onprem and t.onprem[0]["type"] == "VPN"
    hub_e = next(n for n in t.nodes.values() if n.name == "vnet-hub-eastus")
    assert {"Azure Firewall", "Azure Bastion", "VPN gateway"} <= set(hub_e.platform_services)
    assert not next(n for n in t.nodes.values() if n.name == "vnet-dev-eastus").default_route_to_nva
    d = {x.key: x for x in assess_arch.dimensions}
    assert d["hybrid"].pattern == "Site-to-site VPN to on-premises"
    assert d["landing_zone"].pattern.startswith("Landing-zone style")
    assert d["traffic"].pattern.startswith("Defence in depth")
    assert any("spoke-to-spoke" in k for k in d["network"].considerations)
    assert "network topology is multi-region hub-and-spoke" in arch.headline


@pytest.mark.parametrize("vnets,expected", [
    ([("a", ["b", "c"]), ("b", ["a", "c"]), ("c", ["a", "b"])], "Full mesh peering"),
    ([("a", []), ("b", [])], "Isolated virtual networks (no peering)"),
    ([("a", [])], "Single virtual network"),
    ([("hub", ["s1", "s2"], ["AzureFirewallSubnet"]), ("s1", ["hub"]), ("s2", ["hub"])], "Hub-and-spoke"),
    ([("core", ["x", "y", "z"]), ("x", ["core"]), ("y", ["core"]), ("z", ["core"])], "Hub-and-spoke"),
    ([("a", ["b"]), ("b", ["a", "c"]), ("c", ["b"])], "Ad-hoc peering (no central hub)"),
])
def test_topology_patterns(vnets, expected):
    res = [_vnet(v[0], peers=v[1], subnets=v[2] if len(v) > 2 else ()) for v in vnets]
    assert _arch(res).topology.pattern == expected


def test_topology_vwan_and_no_network():
    vwan = to_resource({"name": "vwan", "type": "Microsoft.Network/virtualWans", "properties": {"x": 1}})
    assert _arch([vwan, _vnet("a")]).topology.pattern.startswith("Azure Virtual WAN")
    web = to_resource({"name": "app", "type": "Microsoft.Web/sites", "properties": {"httpsOnly": True}})
    assert _arch([web]).topology.pattern.startswith("No virtual networks")


def test_topology_without_properties(tmp_path):
    from azure_assessment.analysis import architecture

    p = tmp_path / "portal.csv"
    p.write_text("NAME,TYPE,RESOURCE GROUP,LOCATION,SUBSCRIPTION\n"
                 "vnet-hub,Virtual network,rg,East US,Conn\nvnet-app,Virtual network,rg,East US,Prod\n")
    t = architecture.analyse(assess(loaders.load(p))).topology
    assert t.pattern == "Likely hub-and-spoke (inferred from naming)" and t.confidence == "Low"


def test_reports_include_architecture(tmp_path):
    out = generate_reports(sample.build(), tmp_path, "Contoso")
    html = out["html"].read_text()
    assert 'id="architecture"' in html and "<svg class=\"topology\"" in html and "vnet-dev-sandbox" in html
    d = docx.Document(out["docx"])
    heads = [p.text for p in d.paragraphs if p.style.name.startswith("Heading")]
    assert "3.1 Network topology" in heads and "3.7 Operations and monitoring" in heads
    assert len(d.inline_shapes) >= 5  # topology diagram added to the charts


def test_default_output_is_project_reports_folder(tmp_path, monkeypatch):
    from azure_assessment import cli

    assert cli.REPORTS_DIR.name == "reports" and (cli.REPORTS_DIR.parent / "pyproject.toml").exists()
    monkeypatch.setattr(cli, "REPORTS_DIR", tmp_path / "reports")
    monkeypatch.chdir(tmp_path)
    assert main(["demo", "-c", "Fabrikam Ltd"]) == 0
    assert sorted(p.name for p in (tmp_path / "reports").iterdir()) == [
        "fabrikam-ltd-azure-assessment-findings.json", "fabrikam-ltd-azure-assessment.docx",
        "fabrikam-ltd-azure-assessment.html"]
