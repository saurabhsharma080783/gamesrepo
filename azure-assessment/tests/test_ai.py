"""AI narrative: grounding checks, fallbacks, and an end-to-end run against a mock
OpenAI-compatible server (the same API Ollama, llama.cpp and vLLM expose)."""
import json
import re
import threading
from http.server import BaseHTTPRequestHandler, HTTPServer

import docx
import pytest

from azure_assessment.ai import OpenAICompatibleClient, draft, guard
from azure_assessment.analysis.assessor import assess
from azure_assessment.cli import main
from azure_assessment.loaders import sample

FACTS = "Scope: 125 resources in 3 subscription(s). Overall posture score: 62/100. SEC-001 affects 4 of 10."


def test_guard_accepts_grounded_text():
    paras = guard.clean("The estate holds 125 resources across 3 subscriptions and scores 62/100, a result "
                        "that leaves clear room for improvement. SEC-001 affects 4 of the 10 storage accounts "
                        "evaluated, which exposes data to interception in transit.")
    assert guard.check(paras, FACTS) is None


@pytest.mark.parametrize("text, problem", [
    ("The estate holds 130 resources and the posture is weak, with several gaps in security, resilience "
     "and governance that the organisation should address soon.", "130"),
    ("SEC-009 is the main risk in this estate and it should be addressed before anything else, because it "
     "affects production workloads and customer data.", "SEC-009"),
    ("Remediation would save about $1,200 a month across the estate, which the finance team should track "
     "closely alongside the security work described in this report.", "1200"),
    ("Too short.", "too short"),
])
def test_guard_rejects_ungrounded_text(text, problem):
    assert problem in guard.check(guard.clean(text), FACTS)


def test_guard_cleans_markdown_and_preamble():
    paras = guard.clean("Here is the summary:\n\n## Overview\n**Bold** text here.\n\n- first point\n- second")
    assert paras == ["Overview Bold text here.", "first point second"]
    # plain-text headings, as Qwen 2.5 7B produced in the executive summary dry run
    paras = guard.clean("Executive Summary\n\nThe estate is sound.\n\nWhat Is Working Well\n\nTagging is good.")
    assert paras == ["The estate is sound.", "Tagging is good."]
    assert guard.clean("Set `allowBlobPublicAccess` to false.") == ["Set allowBlobPublicAccess to false."]


def test_executive_facts_have_no_rule_ids():
    from azure_assessment.ai import facts
    from azure_assessment.analysis import architecture
    a = assess(sample.build(), customer="Contoso")
    sheet = facts.executive(a, architecture.analyse(a))
    assert "Storage account allows anonymous blob access" in sheet and not guard.RULE_ID.search(sheet)


class _Scripted:
    """Chat client returning scripted replies, recording prompts."""
    model = "test-model"

    def __init__(self, replies):
        self.replies, self.calls = list(replies), []

    def chat(self, system, user):
        self.calls.append(user)
        return self.replies.pop(0) if self.replies else ""


def _grounded_reply(user: str) -> str:
    """Build a reply from statements in the facts, so every figure is grounded."""
    facts = user.split("Facts:", 1)[-1]
    stmts = [ln[2:] for ln in facts.splitlines() if ln.startswith("- ")][:3]
    return ("This area was reviewed from the inventory export and interpreted for the business. "
            + " ".join(stmts) + "\n\nThe observations above should be read together with the related findings.")


def test_draft_retries_then_falls_back():
    a = assess(sample.build(), customer="Contoso")
    bad = "The estate has 99999 resources, far more than expected, which makes it very hard to manage well."
    client = _Scripted([bad, bad])  # executive summary: first draft and retry both invent a figure
    d = draft(a, client, workers=1, retries=1)
    ex = d.parts[0]
    assert ex.kind == "executive" and ex.status == "standard" and ex.attempts == 2 and "99999" in ex.reason
    assert "previous draft was rejected" in client.calls[1]
    assert d.executive() == [] and d.drafted == 0  # later parts got empty replies -> standard wording


class _MockLLM(BaseHTTPRequestHandler):
    def log_message(self, *args):
        pass

    def _send(self, body):
        data = json.dumps(body).encode()
        self.send_response(200)
        self.send_header("Content-Type", "application/json")
        self.send_header("Content-Length", str(len(data)))
        self.end_headers()
        self.wfile.write(data)

    def do_GET(self):
        self._send({"data": [{"id": "mock-7b"}]})

    def do_POST(self):
        req = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
        user = req["messages"][-1]["content"]
        reply = _grounded_reply(user)
        if '"Regional footprint"' in user:  # one section invents a figure every time -> standard wording kept
            reply = "Resources sit in 42 regions. " + reply
        self._send({"choices": [{"message": {"role": "assistant", "content": reply}}]})


@pytest.fixture
def llm_server():
    srv = HTTPServer(("127.0.0.1", 0), _MockLLM)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    yield f"http://127.0.0.1:{srv.server_port}/v1"
    srv.shutdown()


def test_cli_ai_end_to_end(tmp_path, llm_server, capsys):
    assert main(["demo", "-c", "Contoso", "-o", str(tmp_path), "--ai", "--llm-url", llm_server,
                 "--llm-model", "mock-7b"]) == 0
    out = capsys.readouterr().out
    assert "Drafting narrative with mock-7b" in out and "standard Regional footprint" in out

    saved = json.loads((tmp_path / "contoso-azure-assessment-ai-draft.json").read_text())
    by_title = {p["title"]: p for p in saved["parts"]}
    assert saved["model"] == "mock-7b" and saved["drafted"] == saved["total"] - 1
    assert by_title["Regional footprint"]["status"] == "standard" and "42" in by_title["Regional footprint"]["reason"]
    assert by_title["Executive summary"]["status"] == "ai" and by_title["Executive summary"]["facts"]

    html = (tmp_path / "contoso-azure-assessment.html").read_text()
    assert "Use of AI in this report" in html and "mock-7b" in html
    assert html.count('class="ai-mark"') == saved["drafted"]
    assert "interpreted for the business" in html and "42 regions" not in html

    d = docx.Document(tmp_path / "contoso-azure-assessment.docx")
    text = "\n".join(p.text for p in d.paragraphs)
    assert "Use of AI in this report" in text and "interpreted for the business" in text
    assert "42 regions" not in text


def test_cli_ai_server_unreachable_keeps_standard_report(tmp_path, capsys):
    assert main(["demo", "-o", str(tmp_path), "-f", "html", "--ai", "--llm-url", "http://127.0.0.1:9/v1"]) == 0
    err = capsys.readouterr().err
    assert "AI narrative skipped" in err
    html = next(tmp_path.glob("*.html")).read_text()
    assert "Use of AI" not in html and not list(tmp_path.glob("*-ai-draft.json"))


def test_client_reports_missing_model(llm_server):
    with pytest.raises(Exception, match=re.escape("ollama pull other")):
        OpenAICompatibleClient(llm_server, "other").check()


def test_fact_sheets_leave_out_per_resource_rows():
    # A real 7B model misattributed settings when given table rows; the sheets carry summaries only.
    from azure_assessment.ai import facts
    from azure_assessment.analysis import architecture, narrative
    a = assess(sample.build(), customer="Contoso")
    storage = next(s for s in narrative.build(a) if s.key == "storage")
    sheet = facts.section(storage)
    assert "storage accounts" in sheet and "Related findings:" in sheet
    assert not any(str(row[0]) in sheet for _, t in storage.tables for row in t.rows)
    assert "Table '" not in facts.dimension(architecture.analyse(a).dimensions[0])
