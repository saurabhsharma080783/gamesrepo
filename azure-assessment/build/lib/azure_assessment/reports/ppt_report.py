"""PowerPoint (.pptx) executive briefing deck, 16:9."""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Emu, Inches, Pt

from ..analysis.assessor import Assessment
from .theme import ACCENT, INK, INK_2, PRIMARY, SEVERITY_COLORS, SURFACE, hex_to_rgb, rating_color

W, H = Inches(13.333), Inches(7.5)
WHITE = "#FFFFFF"


def _rgb(h):
    return RGBColor(*hex_to_rgb(h))


def _text(slide, x, y, w, h, text, size=14, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    lines = text if isinstance(text, list) else [text]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = str(line)
        r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(size), bold, _rgb(color), "Calibri"
    return tb


def _rect(slide, x, y, w, h, fill, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = _rgb(fill)
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def _slide(prs, title, subtitle=None):
    s = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    _rect(s, 0, 0, W, Inches(0.12), ACCENT)
    _text(s, Inches(0.6), Inches(0.35), Inches(12), Inches(0.7), title, size=28, color=PRIMARY, bold=True)
    if subtitle:
        _text(s, Inches(0.6), Inches(0.98), Inches(12), Inches(0.4), subtitle, size=13, color=INK_2)
    return s


def _picture(slide, path, x, y, max_w, max_h):
    """Place an image scaled to fit inside the box, centred horizontally, keeping its aspect ratio."""
    from PIL import Image

    with Image.open(path) as im:
        iw, ih = im.size
    scale = min(max_w / iw, max_h / ih)
    w, h = int(iw * scale), int(ih * scale)
    return slide.shapes.add_picture(str(path), x + (max_w - w) // 2, y, w, h)


def _footer(s, a: Assessment, n: int):
    _text(s, Inches(0.6), Inches(7.05), Inches(9), Inches(0.3),
          f"{a.customer} · Azure Environment Assessment · Confidential", size=9, color=INK_2)
    _text(s, Inches(11.9), Inches(7.05), Inches(0.9), Inches(0.3), str(n), size=9, color=INK_2, align=PP_ALIGN.RIGHT)


def _kpi(s, x, y, w, value, label, color=PRIMARY):
    _rect(s, x, y, w, Inches(1.25), SURFACE, MSO_SHAPE.ROUNDED_RECTANGLE).adjustments[0] = 0.08
    _text(s, x, y + Inches(0.12), w, Inches(0.7), str(value), size=30, color=color, bold=True, align=PP_ALIGN.CENTER)
    _text(s, x, y + Inches(0.8), w, Inches(0.35), label, size=11, color=INK_2, align=PP_ALIGN.CENTER)


def _table(s, x, y, w, headers, rows, col_widths=None, size=11):
    shape = s.shapes.add_table(len(rows) + 1, len(headers), x, y, w, Inches(0.4) * (len(rows) + 1))
    t = shape.table
    if col_widths:
        total = sum(col_widths)
        for i, cw in enumerate(col_widths):
            t.columns[i].width = Emu(int(w * cw / total))
    for i, h in enumerate(headers):
        c = t.cell(0, i)
        c.fill.solid()
        c.fill.fore_color.rgb = _rgb(PRIMARY)
        c.text = h
        f = c.text_frame.paragraphs[0].runs[0].font
        f.size, f.bold, f.color.rgb = Pt(size), True, _rgb(WHITE)
    for ri, row in enumerate(rows, 1):
        for ci, v in enumerate(row):
            c = t.cell(ri, ci)
            c.fill.solid()
            c.fill.fore_color.rgb = _rgb("#FFFFFF" if ri % 2 else "#F2F5F9")
            c.text = str(v)
            f = c.text_frame.paragraphs[0].runs[0].font
            f.size, f.color.rgb = Pt(size - 1), _rgb(INK)
            if headers[ci] == "Severity" and v in SEVERITY_COLORS:
                f.bold = True
                f.color.rgb = _rgb(SEVERITY_COLORS[v] if v != "Low" else "#9A6700")
    return t


def build(a: Assessment, charts: dict[str, Path], out: Path) -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    n = 0

    # 1. Title
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _rect(s, 0, 0, W, H, PRIMARY)
    _rect(s, Inches(0.6), Inches(2.3), Inches(0.12), Inches(2.4), ACCENT)
    _text(s, Inches(1.0), Inches(2.2), Inches(11), Inches(0.5), "AZURE ENVIRONMENT ASSESSMENT", size=16,
          color="#9EC5F4", bold=True)
    _text(s, Inches(1.0), Inches(2.75), Inches(11), Inches(1.2), a.customer, size=48, color=WHITE, bold=True)
    _text(s, Inches(1.0), Inches(3.9), Inches(11), Inches(0.6),
          "Inventory, Well-Architected findings and remediation roadmap", size=20, color="#CDE2FB")
    _text(s, Inches(1.0), Inches(6.3), Inches(11), Inches(0.4), datetime.now().strftime("%B %Y"), size=14,
          color="#CDE2FB")
    n += 1

    # 2. Agenda
    s = _slide(prs, "Agenda")
    items = ["Executive summary", "Scope & methodology", "Inventory overview", "Assessment scorecard",
             "Key findings", "Remediation roadmap", "Next steps"]
    for i, it in enumerate(items):
        y = Inches(1.6 + i * 0.72)
        _rect(s, Inches(0.8), y, Inches(0.5), Inches(0.5), ACCENT, MSO_SHAPE.OVAL)
        _text(s, Inches(0.8), y, Inches(0.5), Inches(0.5), str(i + 1), size=14, color=WHITE, bold=True,
              align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        _text(s, Inches(1.55), y, Inches(8), Inches(0.5), it, size=20, color=INK, anchor=MSO_ANCHOR.MIDDLE)
    n += 1
    _footer(s, a, n)

    # 3. Executive summary
    s = _slide(prs, "Executive summary",
               f"Overall posture score {a.overall_score}/100 — {a.overall_rating}"
               + (f" ({a.score_basis.lower()})" if a.score_basis else "") if a.overall_score is not None
               else "Overall posture score not available — see data coverage")
    sev = a.severity_counts
    kpis = [(f"{a.total_resources:,}", "Resources", PRIMARY), (len(a.inventory.subscriptions), "Subscriptions", PRIMARY),
            (len(a.by_location()), "Regions", PRIMARY), (a.overall_score_label, "Posture score", rating_color(a.overall_score)),
            (sev["High"], "High findings", SEVERITY_COLORS["High"]), (a.tag_coverage_label, "Tag coverage", PRIMARY)]
    for i, (v, l, c) in enumerate(kpis):
        _kpi(s, Inches(0.6 + i * 2.05), Inches(1.6), Inches(1.9), v, l, c)
    obs = a.key_observations()
    tb = _text(s, Inches(0.6), Inches(3.2), Inches(12.1), Inches(3.6), [f"•  {o}" for o in obs], size=15)
    for p in tb.text_frame.paragraphs:
        p.space_after = Pt(8)
    n += 1
    _footer(s, a, n)

    # 4. Scope & methodology
    s = _slide(prs, "Scope & methodology")
    _text(s, Inches(0.6), Inches(1.5), Inches(5.8), Inches(5), [
        f"•  Based on the inventory export supplied by {a.customer}",
        "•  Every resource evaluated against rules aligned to the Azure Well-Architected Framework",
        "•  Pillars: Security, Reliability, Cost Optimization, Operational Excellence, Governance",
        "•  Scores = severity-weighted pass rate (High 5 · Medium 3 · Low 1)",
        "•  Rating: Good ≥ 80 · Fair 60–79 · Poor < 60",
    ], size=15)
    subs = [(name, sum(1 for r in a.inventory.resources if r.subscription_id == sid))
            for sid, name in a.inventory.subscriptions.items()][:10]
    _table(s, Inches(6.8), Inches(1.6), Inches(5.9), ["Subscription", "Resources"], subs, [4, 1.3], size=12)
    n += 1
    _footer(s, a, n)

    if a.coverage_gaps:
        s = _slide(prs, "Data coverage & limitations",
                   "Checks below lacked data in the inventory export and are excluded from scores")
        gaps = a.coverage_gaps[:9]
        _table(s, Inches(0.6), Inches(1.55), Inches(12.1), ["Rule", "Check", "Pillar", "Data needed", "Not assessed"],
               [(g.rule.id, g.rule.title, g.rule.pillar, g.missing, f"{g.skipped} of {g.applicable}") for g in gaps],
               [1, 5, 2.4, 1.6, 1.5], size=12)
        more = f"  (+{len(a.coverage_gaps) - 9} more in the Word report)" if len(a.coverage_gaps) > 9 else ""
        _text(s, Inches(0.6), Inches(6.35), Inches(12.1), Inches(0.5),
              "For a complete assessment, share a Resource Graph export including tags, sku, zones and properties."
              + more, size=12, color=INK_2)
        n += 1
        _footer(s, a, n)

    # 5. Inventory overview
    s = _slide(prs, "Inventory overview", f"{a.total_resources:,} resources · {a.resource_group_count} resource groups")
    _picture(s, charts["by_type"], Inches(0.5), Inches(1.5), Inches(6.1), Inches(5.4))
    _picture(s, charts["by_region"], Inches(6.8), Inches(1.5), Inches(6.1), Inches(5.4))
    n += 1
    _footer(s, a, n)

    # 6. Scorecard
    s = _slide(prs, "Assessment scorecard", "Severity-weighted pass rate per pillar")
    _picture(s, charts["pillars"], Inches(0.5), Inches(1.6), Inches(6.2), Inches(5.2))
    _table(s, Inches(7.0), Inches(1.7), Inches(5.8), ["Pillar", "Score", "Rating", "High", "Med", "Low"],
           [(p.pillar, p.score_label, p.rating, p.high, p.medium, p.low) for p in a.pillar_scores],
           [2.5, 1.05, 1.15, 0.85, 0.8, 0.8], size=12)
    n += 1
    _footer(s, a, n)

    # 7. Findings distribution
    s = _slide(prs, "Findings distribution", f"{len(a.findings)} findings across {a.affected_resources} resources")
    _picture(s, charts["severity"], Inches(0.5), Inches(1.7), Inches(5.8), Inches(5.1))
    _picture(s, charts["pillar_findings"], Inches(6.6), Inches(1.7), Inches(6.3), Inches(5.1))
    n += 1
    _footer(s, a, n)

    # 8+. Top findings (paged, 8 per slide)
    top = a.top_rules(16)
    for page in range(0, max(len(top), 1), 8):
        chunk = top[page:page + 8]
        s = _slide(prs, "Key findings" + (" (cont.)" if page else ""), "Ranked by severity, then number of resources affected")
        if chunk:
            _table(s, Inches(0.6), Inches(1.55), Inches(12.1), ["ID", "Finding", "Severity", "Pillar", "Affected"],
                   [(r.rule.id, r.rule.title, r.rule.severity, r.rule.pillar, r.affected) for r in chunk],
                   [1, 5.5, 1.2, 2.4, 1.1], size=13)
        else:
            _text(s, Inches(0.6), Inches(2), Inches(12), Inches(1), "No findings were raised.", size=18)
        n += 1
        _footer(s, a, n)

    # Roadmap
    s = _slide(prs, "Remediation roadmap", "Sequenced by risk and effort")
    plan = a.roadmap()
    col_w = Inches(3.95)
    for i, (phase, items) in enumerate(plan.items()):
        x = Inches(0.6) + i * (col_w + Inches(0.15))
        _rect(s, x, Inches(1.55), col_w, Inches(0.6), [ACCENT, "#256ABF", PRIMARY][i])
        _text(s, x + Inches(0.15), Inches(1.55), col_w - Inches(0.3), Inches(0.6), phase, size=13, color=WHITE,
              bold=True, anchor=MSO_ANCHOR.MIDDLE)
        _rect(s, x, Inches(2.15), col_w, Inches(4.7), SURFACE)
        lines = [f"•  {it.rule.title} ({it.affected})" for it in items[:9]] or ["No actions"]
        if len(items) > 9:
            lines.append(f"…and {len(items) - 9} more")
        tb = _text(s, x + Inches(0.15), Inches(2.3), col_w - Inches(0.3), Inches(4.4), lines, size=12)
        for p in tb.text_frame.paragraphs:
            p.space_after = Pt(6)
    n += 1
    _footer(s, a, n)

    # Next steps
    s = _slide(prs, "Recommended next steps")
    steps = [
        "Remediate high-severity security exposures (public access, HTTP, open management ports)",
        "Reclaim idle spend: orphaned disks, unused public IPs, stopped VMs and empty plans",
        "Enforce tagging and allowed regions through Azure Policy at management-group scope",
        "Design zone-redundant patterns for production workloads",
        "Re-run this assessment monthly to track posture trend",
    ]
    tb = _text(s, Inches(0.6), Inches(1.6), Inches(12), Inches(5), [f"{i + 1}.  {t}" for i, t in enumerate(steps)],
               size=18)
    for p in tb.text_frame.paragraphs:
        p.space_after = Pt(14)
    n += 1
    _footer(s, a, n)

    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    return out
